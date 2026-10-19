"""文档加载与解析

职责：目录扫描 -> 单文件解析 -> 文本抽取 -> 清洗去重 -> 返回结构化文档
不负责：切分、向量化、检索（见 core/knowledge_base.py）

支持格式：
  文本   .txt .md .markdown .log .ini .cfg .rst .csv .tsv .json
  文档   .pdf .docx .pptx .xlsx .xlsm
  网页   .html .htm
  图片   .png .jpg .jpeg .jp2 .webp .gif .bmp .tif .tiff（OCR）
  旧格式 .doc .ppt .xls（本地库无解，需 MinerU 云端）

PDF/图片/扫描件解析使用 MinerU 云端 API v4，未配置 Token 时自动跳过。
所有第三方解析库均为可选依赖，缺失时跳过对应格式，不会导致启动失败。
"""
import os
import io
import sys
import csv
import json
import time
import codecs
import hashlib
import zipfile
import tempfile
from pathlib import Path

MINERU_API_BASE = "https://mineru.net/api/v4"

# 项目根：core/xxx.py -> core -> backend -> 项目根。数据目录一律基于 __file__ 推算，
# 既不依赖启动目录，也不硬编码绝对路径。
_BACKEND_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_PROJECT_DIR = os.path.dirname(_BACKEND_DIR)
DEFAULT_DATA_DIR = os.path.join(_PROJECT_DIR, "data", "raw")


def _force_utf8_streams():
    """Windows 控制台默认 GBK，输出 emoji / 特殊字符会抛 UnicodeEncodeError"""
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass


_force_utf8_streams()

# 文本类扩展名
TEXT_EXT = {".txt", ".md", ".markdown", ".log", ".ini", ".cfg", ".rst"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".jp2", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
MINERU_EXT = {".pdf", ".doc", ".docx", ".ppt", ".pptx"} | IMAGE_EXT


class MinerUError(Exception):
    """MinerU 云端解析失败"""


class MinerUCloudClient:
    """MinerU 云端解析（API v4）：申请上传 URL -> PUT -> 轮询 -> 下载 zip -> 取 full.md"""

    def __init__(self, token=None, model_version="pipeline", language="ch",
                 cache_dir=None, poll_interval=3.0, max_wait=300, timeout=60):
        self.token = token or os.getenv("MINERU_API_TOKEN") or os.getenv("MINERU_API_KEY")
        self.model_version = model_version
        self.language = language
        self.poll_interval = poll_interval
        self.max_wait = max_wait
        self.timeout = timeout
        self.cache_dir = cache_dir or os.path.join(
            tempfile.gettempdir(), "tour_guide_mineru_cache")
        self._session = None

    @property
    def available(self):
        return bool(self.token)

    def _get_session(self):
        if self._session is None:
            import requests  # 可选依赖
            self._session = requests.Session()
        return self._session

    def _request(self, method, url, **kw):
        """发起请求并校验：非 200 或 code!=0 一律抛错，成功返回 data"""
        resp = getattr(self._get_session(), method)(url, timeout=self.timeout, **kw)
        if resp.status_code != 200:
            raise MinerUError(f"HTTP {resp.status_code}: {resp.text[:200]}")
        body = resp.json()
        if body.get("code") != 0:
            raise MinerUError(body.get("msg") or "接口返回错误")
        return body.get("data") or {}

    def _cache_path(self, filepath):
        """同名同大小同修改时间视为同一份文件，复用解析结果省额度"""
        st = os.stat(filepath)
        raw = f"{os.path.basename(filepath)}|{st.st_size}|{int(st.st_mtime)}|{self.model_version}"
        digest = hashlib.md5(raw.encode("utf-8")).hexdigest()
        return os.path.join(self.cache_dir, digest + ".md")

    def parse(self, filepath, is_ocr=True):
        """解析单个文件，返回 Markdown 文本"""
        if not self.available:
            raise MinerUError("未配置 MINERU_API_TOKEN")
        if not os.path.exists(filepath):
            raise MinerUError(f"文件不存在: {filepath}")

        cache = self._cache_path(filepath)
        if os.path.exists(cache):
            print(f"    ♻️  MinerU 缓存命中: {os.path.basename(filepath)}")
            with open(cache, "r", encoding="utf-8") as f:
                return f.read()

        headers = {"Authorization": f"Bearer {self.token}"}

        # 1. 申请上传 URL
        data = self._request("post", f"{MINERU_API_BASE}/file-urls/batch",
                             headers={**headers, "Content-Type": "application/json"},
                             json={"files": [{"name": os.path.basename(filepath),
                                              "is_ocr": bool(is_ocr)}],
                                   "model_version": self.model_version,
                                   "language": self.language})
        batch_id, upload_url = data["batch_id"], data["file_urls"][0]

        # 2. 上传文件（不带 Content-Type）
        with open(filepath, "rb") as f:
            up = self._get_session().put(upload_url, data=f, timeout=self.timeout)
        if up.status_code != 200:
            raise MinerUError(f"文件上传失败 HTTP {up.status_code}")

        # 3. 轮询结果
        result_url = f"{MINERU_API_BASE}/extract-results/batch/{batch_id}"
        deadline = time.time() + self.max_wait
        while time.time() < deadline:
            time.sleep(self.poll_interval)
            items = self._request("get", result_url, headers=headers).get("extract_result")
            if not items:
                continue
            item = items[0]
            state = item.get("state")
            if state == "done":
                text = self._download_markdown(item.get("full_zip_url"))
                os.makedirs(self.cache_dir, exist_ok=True)
                with open(cache, "w", encoding="utf-8") as f:
                    f.write(text)
                return text
            if state == "failed":
                raise MinerUError(f"云解析失败: {item.get('err_msg') or '未知原因'}")
            prog = item.get("extract_progress") or {}
            if prog.get("total_pages"):
                print(f"    ⏳ 云解析中 {prog.get('extracted_pages', 0)}/{prog['total_pages']} 页")

        raise MinerUError(f"云解析超时（{self.max_wait}s）")

    def _download_markdown(self, zip_url):
        """下载结果包，优先取 full.md"""
        if not zip_url:
            raise MinerUError("解析完成但未返回结果包地址")
        resp = self._get_session().get(zip_url, timeout=self.timeout)
        if resp.status_code != 200:
            raise MinerUError(f"结果包下载失败 HTTP {resp.status_code}")
        with zipfile.ZipFile(io.BytesIO(resp.content)) as zf:
            names = [n for n in zf.namelist() if n.lower().endswith(".md")]
            target = next((n for n in names if os.path.basename(n).lower() == "full.md"), None)
            target = target or (names[0] if names else None)
            if not target:
                raise MinerUError("结果包中未找到 Markdown 文件")
            return zf.read(target).decode("utf-8", errors="replace")


class DocumentLoader:
    """多格式文档加载器，与向量存储完全解耦"""

    def __init__(self, data_dir=DEFAULT_DATA_DIR, mineru_token=None,
                 mineru_model="pipeline", pdf_parser="auto",
                 pdf_text_threshold=50, max_file_mb=200, max_pdf_pages=600,
                 min_content_length=10, recursive=False, enable_dedup=True):
        self.data_dir = data_dir
        self.pdf_parser = pdf_parser                      # auto | pypdf | mineru
        self.pdf_text_threshold = pdf_text_threshold      # 平均每页字符数低于此值判定为扫描件
        self.max_file_mb = max_file_mb
        self.max_pdf_pages = max_pdf_pages
        self.min_content_length = min_content_length
        self.recursive = recursive
        self.enable_dedup = enable_dedup

        self.mineru = MinerUCloudClient(token=mineru_token, model_version=mineru_model)
        self.stats = {}                                   # 最近一次加载统计

    # ==================== 编码探测 ====================
    def _detect_encoding(self, filepath):
        with open(filepath, "rb") as f:
            raw = f.read(8192)
        if raw.startswith(codecs.BOM_UTF8):
            return "utf-8-sig"
        if raw.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
            return "utf-16"
        for enc in ("utf-8", "gbk", "gb18030", "big5"):
            try:
                raw.decode(enc)
                return enc
            except (UnicodeDecodeError, LookupError):
                continue
        try:
            from charset_normalizer import from_bytes
            best = from_bytes(raw).best()
            if best and best.encoding:
                return best.encoding
        except Exception:
            pass
        return "utf-8"

    def _read_text_file(self, filepath):
        enc = self._detect_encoding(filepath)
        with open(filepath, "r", encoding=enc, errors="replace") as f:
            return f.read()

    # ==================== 各格式解析 ====================
    def _read_csv(self, filepath):
        enc = self._detect_encoding(filepath)
        rows = []
        with open(filepath, "r", encoding=enc, errors="replace", newline="") as f:
            sample = f.read(8192)
            f.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            except Exception:
                dialect = csv.excel
            reader = csv.reader(f, dialect)
            for row in reader:
                cells = [c.strip() for c in row if c and c.strip()]
                if cells:
                    rows.append(" | ".join(cells))
        return "\n".join(rows)

    def _read_json(self, filepath):
        def flat(obj, prefix=""):
            lines = []
            if isinstance(obj, dict):
                for k, v in obj.items():
                    lines.extend(flat(v, f"{prefix}{k}." if prefix else f"{k}."))
            elif isinstance(obj, list):
                for i, v in enumerate(obj):
                    lines.extend(flat(v, f"{prefix}[{i}]."))
            else:
                key = prefix.rstrip(".")
                lines.append(f"{key}: {obj}" if key else str(obj))
            return lines

        with open(filepath, "r", encoding="utf-8", errors="replace") as f:
            data = json.load(f)
        return "\n".join(flat(data))

    def _read_html(self, filepath):
        try:
            from bs4 import BeautifulSoup
        except ImportError:
            raise RuntimeError("未安装 beautifulsoup4，无法解析 HTML")
        raw = self._read_text_file(filepath)
        soup = BeautifulSoup(raw, "html.parser")
        for tag in soup(["script", "style", "noscript", "nav", "footer"]):
            tag.decompose()
        text = soup.get_text("\n")
        return "\n".join(line.strip() for line in text.splitlines() if line.strip())

    def _read_docx(self, filepath):
        try:
            from docx import Document
        except ImportError:
            raise RuntimeError("未安装 python-docx，无法解析 DOCX")
        doc = Document(filepath)
        parts = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n\n".join(parts)

    def _read_pptx(self, filepath):
        try:
            from pptx import Presentation
        except ImportError:
            raise RuntimeError("未安装 python-pptx，无法解析 PPTX")
        prs = Presentation(filepath)
        parts = []
        for idx, slide in enumerate(prs.slides, 1):
            lines = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        t = para.text.strip()
                        if t:
                            lines.append(t)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        cells = [c.text.strip() for c in row.cells if c.text.strip()]
                        if cells:
                            lines.append(" | ".join(cells))
            if lines:
                parts.append(f"[第 {idx} 页]\n" + "\n".join(lines))
        return "\n\n".join(parts)

    def _read_xlsx(self, filepath):
        try:
            import openpyxl
        except ImportError:
            raise RuntimeError("未安装 openpyxl，无法解析 Excel")
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        parts = []
        try:
            for ws in wb.worksheets:
                rows = []
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i >= 5000:
                        rows.append("... (表格内容过长已截断)")
                        break
                    cells = [str(c).strip() for c in row
                             if c is not None and str(c).strip()]
                    if cells:
                        rows.append(" | ".join(cells))
                if rows:
                    parts.append(f"[工作表 {ws.title}]\n" + "\n".join(rows))
        finally:
            wb.close()
        return "\n\n".join(parts)

    def _read_pdf_pypdf(self, filepath):
        try:
            from pypdf import PdfReader
        except ImportError:
            raise RuntimeError("未安装 pypdf，无法解析 PDF")
        reader = PdfReader(filepath)
        total = len(reader.pages)
        limit = min(total, self.max_pdf_pages)
        pages = []
        for i in range(limit):
            try:
                pages.append(reader.pages[i].extract_text() or "")
            except Exception:
                pages.append("")
        if total > limit:
            print(f"    ⚠️  PDF 共 {total} 页，仅解析前 {limit} 页")
        return "\n".join(pages), limit

    def _read_pdf(self, filepath):
        """PDF 解析：auto 模式下先用 pypdf，识别为扫描件再走 MinerU 云端"""
        if self.pdf_parser == "mineru":
            return self._read_by_mineru(filepath), "mineru"

        if self.pdf_parser == "pypdf":
            text, _ = self._read_pdf_pypdf(filepath)
            return text, "pypdf"

        # auto：优先本地解析，失败或判定为扫描件再走云端
        try:
            text, pages = self._read_pdf_pypdf(filepath)
        except Exception as e:
            print(f"    ⚠️  本地解析失败（{e}），改用 MinerU 云端")
            return self._read_by_mineru(filepath), "mineru"

        chars = len(text.strip())
        avg = chars / pages if pages else 0
        if pages and avg >= self.pdf_text_threshold and chars >= 100:
            return text, "pypdf"
        print(f"    🔍 文本层稀疏（{chars} 字 / {pages} 页），判定为扫描件，改用 MinerU 云端")
        return self._read_by_mineru(filepath), "mineru"

    def _read_image(self, filepath):
        """图片 OCR：统一交给 MinerU 云端；tiff 先转 png"""
        target, tmp_path = filepath, None
        if Path(filepath).suffix.lower() in (".tif", ".tiff"):
            try:
                from PIL import Image
                tmp_path = os.path.join(tempfile.gettempdir(),
                                        f"kb_img_{os.getpid()}_{int(time.time())}.png")
                with Image.open(filepath) as im:
                    im.convert("RGB").save(tmp_path, "PNG")
                target = tmp_path
            except Exception as e:
                print(f"    ⚠️  TIFF 转换失败，尝试直接上传: {e}")
        try:
            return self._read_by_mineru(target)
        finally:
            if tmp_path and os.path.exists(tmp_path):
                try:
                    os.remove(tmp_path)
                except OSError:
                    pass

    def _read_by_mineru(self, filepath, is_ocr=True):
        if not self.mineru.available:
            raise MinerUError("未配置 MINERU_API_TOKEN，无法使用云解析")
        return self.mineru.parse(filepath, is_ocr=is_ocr)

    # ==================== 单文件调度 ====================
    def _parse_file(self, filepath):
        """解析单个文件，返回 (content, parser_name)；失败抛出 Exception"""
        ext = Path(filepath).suffix.lower()
        size_mb = os.path.getsize(filepath) / 1024 / 1024
        if size_mb > self.max_file_mb:
            raise RuntimeError(f"文件过大 {size_mb:.1f}MB > {self.max_file_mb}MB")

        if ext in TEXT_EXT:
            return self._read_text_file(filepath), "text"
        if ext in (".csv", ".tsv"):
            return self._read_csv(filepath), "csv"
        if ext == ".json":
            return self._read_json(filepath), "json"
        if ext in (".html", ".htm"):
            return self._read_html(filepath), "html"
        if ext == ".docx":
            return self._read_docx(filepath), "docx"
        if ext == ".pptx":
            return self._read_pptx(filepath), "pptx"
        if ext in (".xlsx", ".xlsm"):
            return self._read_xlsx(filepath), "xlsx"
        if ext == ".pdf":
            return self._read_pdf(filepath)
        if ext in IMAGE_EXT:
            return self._read_image(filepath), "mineru-ocr"
        if ext in (".doc", ".ppt", ".xls"):
            # 旧二进制格式，本地库无解，交给云端
            return self._read_by_mineru(filepath), "mineru"
        raise RuntimeError(f"暂不支持的格式: {ext}")

    def _iter_files(self):
        if not os.path.isdir(self.data_dir):
            print(f"⚠️  知识库目录不存在: {self.data_dir}")
            return []
        files = []
        if self.recursive:
            for root, dirs, names in os.walk(self.data_dir):
                dirs[:] = [d for d in dirs if not d.startswith(".")]
                files.extend(os.path.join(root, n)
                             for n in names if not n.startswith("."))
        else:
            for name in sorted(os.listdir(self.data_dir)):
                path = os.path.join(self.data_dir, name)
                if os.path.isfile(path) and not name.startswith("."):
                    files.append(path)
        return sorted(files)

    # ==================== 文档加载 ====================
    def load_documents(self):
        docs, seen = [], set()
        failed, skipped = [], []
        files = self._iter_files()
        print(f"📂 扫描目录: {self.data_dir}（{len(files)} 个文件）")

        for filepath in files:
            filename = os.path.basename(filepath)
            try:
                content, parser = self._parse_file(filepath)
            except Exception as e:
                failed.append((filename, str(e)))
                print(f"  ❌ {filename}: {e}")
                continue

            content = (content or "").strip()
            if len(content) < self.min_content_length:
                skipped.append((filename, f"内容过短 {len(content)} 字"))
                print(f"  ⏭️  {filename}: 内容为空或过短，已跳过")
                continue

            if self.enable_dedup:
                digest = hashlib.sha1(content.encode("utf-8")).hexdigest()
                if digest in seen:
                    skipped.append((filename, "内容重复"))
                    print(f"  ⏭️  {filename}: 内容与已有文档重复，已跳过")
                    continue
                seen.add(digest)

            docs.append({
                "content": content,
                "source": filename,
                "file_type": Path(filepath).suffix.lower().lstrip("."),
                "parser": parser,
                "size": os.path.getsize(filepath),
            })
            print(f"  📄 {filename}: {len(content)} 字 [{parser}]")

        self.stats = {
            "total_files": len(files),
            "loaded": len(docs),
            "failed": failed,
            "skipped": skipped,
        }
        print(f"✅ 共加载 {len(docs)} 个文档"
              + (f"，失败 {len(failed)} 个" if failed else "")
              + (f"，跳过 {len(skipped)} 个" if skipped else ""))
        return docs


if __name__ == "__main__":
    docs = DocumentLoader().load_documents()
    for doc in docs:
        print(f"  {doc['source']} | {doc['parser']} | {len(doc['content'])} 字")
